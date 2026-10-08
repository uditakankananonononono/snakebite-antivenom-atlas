#!/usr/bin/env python3
import pathlib,csv,json,html,collections
import matplotlib;matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Image,PageBreak,Spacer
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=pathlib.Path(__file__).resolve().parent.parent;P=R/'paper'
def read(p):return list(csv.DictReader(open(R/p)))
sections=json.loads((P/'content.json').read_text());gap=read('results/gap_table.csv');sp=read('results/consolidated_species.csv');who=read('data/closure/who_catalogue.csv')
plt.rcParams.update({'font.family':'serif','font.size':10})
f,a=plt.subplots(figsize=(6.5,2.5));a.bar(['Critical','High','Data gap','Covered'],[16,116,38,124],color=['#a64545','#91a6bb','#b4c4d2','#2367a0']);a.set(ylabel='Inventory records',ylim=(0,145))
for i,n in enumerate([16,116,38,124]):a.text(i,n+3,str(n),ha='center')
f.tight_layout();f.savefig(P/'gap.png',dpi=180);plt.close(f)
f,a=plt.subplots(figsize=(6.5,2.5));a.plot([60,70,80,90],[72.8,66.9,55.5,19],marker='o',color='#2367a0');a.set(xlabel='Normalized edit-similarity threshold (%)',ylabel='Passing scored toxins (%)',ylim=(0,100));a.grid(alpha=.2);f.tight_layout();f.savefig(P/'sensitivity.png',dpi=180);plt.close(f)
style=getSampleStyleSheet();style.add(ParagraphStyle(name='body',fontName='Times-Roman',fontSize=11,leading=16,spaceAfter=14));style.add(ParagraphStyle(name='head',fontName='Times-Bold',fontSize=20,leading=24,textColor=colors.HexColor('#174d80'),spaceAfter=18));style.add(ParagraphStyle(name='cell',fontName='Times-Roman',fontSize=9,leading=12));style.add(ParagraphStyle(name='caption',fontName='Times-Italic',fontSize=9,leading=12,spaceAfter=10))
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
 doc.styles[name].font.name='Times New Roman'
border=OxmlElement('w:pgBorders');border.set(qn('w:offsetFrom'),'page')
for side in ['top','left','bottom','right']:
 x=OxmlElement('w:'+side)
 for k,v in [('val','single'),('sz','6'),('space','20'),('color','2367A0')]:x.set(qn('w:'+k),v)
 border.append(x)
sec._sectPr.append(border);story=[];md=[]
for i,s in enumerate(sections):
 if i:story.append(PageBreak());doc.add_page_break()
 story.append(Paragraph(html.escape(s['title']),style['head']));doc.add_heading(s['title'],1);md.append('# '+s['title'])
 for text in s['text']:story.append(Paragraph(html.escape(text),style['body']));doc.add_paragraph(text);md.append(text+'\n')
 if s.get('figure'):
  fn=s['figure']+'.png';story.append(Image(str(P/fn),width=468,height=180));doc.add_picture(str(P/fn),width=Inches(6.3))
  caption='Figure. Historical inventory gap classes; counts are records, not independent species.' if s['figure']=='gap' else 'Figure. Similarity threshold sensitivity; sequence scores are not efficacy.'
  story.append(Paragraph(caption,style['caption']));doc.add_paragraph(caption);md.append('![]('+fn+')\n'+caption)
 tables=[]
 if s.get('table'):tables.append(s['table'])
 if s.get('dynamic')=='critical':tables.append([['Baseline label','Toxins','Region']]+[[r['species'],r['n_toxins'],r['regions'].replace('|', '; ') or 'unmapped'] for r in gap if r['gap_class'].startswith('CRITICAL')])
 if s.get('dynamic')=='collisions':tables.append([['Group label','Rows']]+[[k,str(v)] for k,v in sorted(collections.Counter(r['species'] for r in sp).items()) if v>1])
 if s.get('dynamic')=='who':tables.append([['Name','Manufacturer','Status']]+[[r['name'],r['manufacturer'],'positive' if r['status'].startswith('positive') else 'under assessment'] for r in who[:10]])
 for tab in tables:
  n=len(tab[0]);widths=[468/n]*n
  if s.get('dynamic')=='critical':widths=[170,50,248]
  if s.get('dynamic')=='who':widths=[205,163,100]
  t=Table([[Paragraph(html.escape(str(x)),style['cell']) for x in row] for row in tab],colWidths=widths,repeatRows=1)
  t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6eff7')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#9fb9d0')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]));story.append(t);story.append(Spacer(1,8))
  dt=doc.add_table(rows=0,cols=n);dt.style='Table Grid'
  for row in tab:
   cells=dt.add_row().cells
   for c,text in zip(cells,row):c.text=str(text)
  md.append('\n'.join('| '+' | '.join(map(str,row))+' |' for row in tab))
(P/'B09_audit_paper.md').write_text('\n\n'.join(md)+'\n');doc.save(P/'B09_audit_paper.docx')
def frame(c,d):
 c.setStrokeColor(colors.HexColor('#2367a0'));c.setLineWidth(.8);c.rect(27,27,558,738);c.setFont('Times-Roman',9);c.drawString(54,38,'B09 | Historical evidence audit | Unsealed closeout');c.drawRightString(558,38,str(d.page))
SimpleDocTemplate(str(P/'B09_audit_paper.pdf'),pagesize=(612,792),leftMargin=72,rightMargin=72,topMargin=54,bottomMargin=60,title='B09 snakebite antivenom evidence audit').build(story,onFirstPage=frame,onLaterPages=frame)
